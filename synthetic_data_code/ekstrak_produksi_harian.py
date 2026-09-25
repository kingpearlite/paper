"""
Ekstraksi data harian per unit dari laporan bulanan PLTD Gili Ketapang (LAPORAN PLTD ... .xlsx).

Menulis `load data/produksi_kwh_harian.csv`: satu baris per hari, dengan produksi kWh, stand meter kWh,
BBM, beban Siang/Malam dan jam nyala untuk Cummins 2 dan 3. Produksi kWh (baris 30-31) menggantikan
pembacaan beban Siang/Malam (baris 34-35) sebagai sumber level beban: 91 dari 365 hari pembacaan kW
adalah salinan templat, sedangkan produksi kWh cocok dengan selisih stand meter setiap hari.

Pemeriksaan (skrip berhenti jika gagal):
  1. jumlah hari = 365 (Agustus 2025 - Juli 2026);
  2. produksi harian = selisih stand meter kumulatif (toleransi 1 kWh), untuk kedua unit;
  3. tidak ada blok salinan (>= 4 hari berturut-turut identik dengan hari-hari sebelumnya) di kolom kWh.

Pemakaian (sumber hanya dibaca):
    C:/Users/<user>/miniconda3/python.exe ekstrak_produksi_harian.py ["<workdir>/Microgridpy/Gili Ketapang"]
"""
import calendar
import datetime as dt
import glob
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd

SRC = Path(sys.argv[1] if len(sys.argv) > 1 else r"<workdir>/Microgridpy\Gili Ketapang")
OUT = Path(__file__).resolve().parent / "load data" / "produksi_kwh_harian.csv"
BULAN = {"JANUARI": 1, "FEBRUARI": 2, "MARET": 3, "APRIL": 4, "MEI": 5, "JUNI": 6, "JULI": 7,
         "AGUSTUS": 8, "SEPTEMBER": 9, "OKTOBER": 10, "NOVEMBER": 11, "DESEMBER": 12}
# Nomor baris di lembar pertama setiap laporan; kolom F = tanggal 1, kolom E = hari terakhir bulan lalu.
ROWS = {"fuel_c2_l": 19, "fuel_c3_l": 20, "stand_c2_kwh": 26, "stand_c3_kwh": 27,
        "prod_c2_kwh": 30, "prod_c3_kwh": 31, "siang_kw": 34, "malam_kw": 35,
        "hours_c2": 57, "hours_c3": 58}


def baca_laporan(src):
    rows = []
    for f in sorted(glob.glob(str(src / "LAPORAN PLTD*.xlsx"))):
        ws = openpyxl.load_workbook(f, data_only=True).worksheets[0]
        # Label baris harus di tempat yang diharapkan; kalau tidak, tata letak file ini berbeda.
        assert "Produksi" in str(ws.cell(row=29, column=2).value), f
        assert str(ws.cell(row=57, column=3).value).strip() == "Cummins 2", f
        m = re.search(r"BULAN\s+([A-Z]+)\s+TAHUN\s+(\d{4})", ws.cell(row=2, column=5).value, re.I)
        month, year = BULAN[m.group(1).upper()], int(m.group(2))
        prev_stand = {"c2": ws.cell(row=26, column=5).value, "c3": ws.cell(row=27, column=5).value}
        for d in range(1, calendar.monthrange(year, month)[1] + 1):
            rec = {"date": dt.date(year, month, d), "file": Path(f).name}
            for k, r in ROWS.items():
                rec[k] = ws.cell(row=r, column=5 + d).value
            if d == 1:
                rec["stand_c2_prev"], rec["stand_c3_prev"] = prev_stand["c2"], prev_stand["c3"]
            rows.append(rec)
    df = pd.DataFrame(rows).sort_values("date").reset_index(drop=True)
    num = [c for c in df.columns if c not in ("date", "file")]
    df[num] = df[num].apply(pd.to_numeric, errors="coerce")
    return df


def hari_salinan(df, cols, min_run=4):
    """Hari yang termasuk blok >= min_run hari berturut-turut yang identik dengan blok sebelumnya."""
    vals = [tuple(r) for r in df[cols].itertuples(index=False)]
    n, hit = len(vals), set()
    for i in range(n):
        for j in range(i):
            k = 0
            while i + k < n and j + k < i and vals[i + k] == vals[j + k] and not all(pd.isna(x) for x in vals[i + k]):
                k += 1
            if k >= min_run:
                hit.update(range(i, i + k))
    return sorted(hit)


def main():
    df = baca_laporan(SRC)
    gagal = []
    print(f"{len(df)} hari, {df['date'].min()} s.d. {df['date'].max()}")
    if len(df) != 365:
        gagal.append(f"jumlah hari {len(df)}, bukan 365")
    first = pd.to_datetime(df["date"]).dt.day == 1
    for u in ("c2", "c3"):
        prev = df[f"stand_{u}_kwh"].shift(1).where(~first, df[f"stand_{u}_prev"])
        ok = ((df[f"stand_{u}_kwh"] - prev) - df[f"prod_{u}_kwh"].fillna(0)).abs() <= 1
        print(f"Cummins {u[1]}: produksi = selisih stand meter pada {int(ok.sum())} dari {len(df)} hari")
        if not ok.all():
            gagal.append(f"Cummins {u[1]}: {int((~ok).sum())} hari tidak cocok dengan meter")
    salinan = hari_salinan(df, ["prod_c2_kwh", "prod_c3_kwh"])
    print(f"Blok salinan di kolom produksi kWh: {len(salinan)} hari")
    if salinan:
        gagal.append(f"{len(salinan)} hari salinan di kolom produksi kWh")
    if gagal:
        sys.exit("GAGAL: " + "; ".join(gagal))
    total = (df["prod_c2_kwh"].fillna(0) + df["prod_c3_kwh"].fillna(0)).sum()
    print(f"Produksi total Cummins 2 + 3: {total / 1e3:.1f} MWh")
    df.to_csv(OUT, index=False)
    print("ditulis", OUT)


if __name__ == "__main__":
    main()
