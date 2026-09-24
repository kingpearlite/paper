"""How often does the one-partial-unit dead band force the battery to discharge in a certified design?
Achievable genset output under Eq. 7 (N units, at most one below full): {0} U [kP+mP, (k+1)P], k=0..N-1.
Without storage, an hour is servable iff some achievable g lies in [net, dem] (g >= net = dem-RES; g <= dem by
Eq. 5). Otherwise the battery must supply at least net - max{achievable g <= net}."""
import sys, glob, os
import numpy as np
import openpyxl

folder, N = sys.argv[1], int(sys.argv[2])
P, m = 470.0, 0.25
lo = np.array([k * P + m * P for k in range(N)])   # lower edge of each piece
hi = np.array([(k + 1) * P for k in range(N)])      # upper edge


def forced(net, dem):
    net = np.maximum(net, 0.0)
    ok = (net <= 1e-9) | np.any((lo[None, :] <= dem[:, None] + 1e-9) & (hi[None, :] >= net[:, None] - 1e-9), axis=1)
    below = np.where(hi[None, :] <= net[:, None], hi[None, :], 0.0).max(axis=1)  # best achievable output under net
    return np.where(ok, 0.0, net - below)


files = sorted(glob.glob(os.path.join(folder, "Time_Series_SC_*.xlsx")),
               key=lambda p: int(p.rsplit("_", 1)[1].split(".")[0]))
T = dict(h=0, f=0.0, dis=0.0, hours=0)
for f in files:
    sc = int(f.rsplit("_", 1)[1].split(".")[0])
    wb = openpyxl.load_workbook(f, read_only=True)
    h = fe = dsum = 0.0
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(min_row=5, values_only=True))
        a = np.array([[float(x or 0) for x in r[1:12]] for r in rows if r[0] is not None])
        dem, res, dis = a[:, 0], a[:, 1], a[:, 6]
        fz = forced(dem - res, dem)
        h += (fz > 1e-6).sum(); fe += fz.sum(); dsum += dis.sum()
        T["hours"] += len(dem)
    wb.close()
    T["h"] += h; T["f"] += fe; T["dis"] += dsum
    print(f"SC {sc:>2}: {int(h):6d} hours need the battery to bridge the band, {fe/1e3:8.1f} MWh forced "
          f"({100*fe/dsum:5.2f}% of the {dsum/1e6:.2f} GWh the battery actually discharged)")
print(f"ALL: {int(T['h'])} of {T['hours']} hours, {T['f']/1e3:.1f} MWh forced, {100*T['f']/T['dis']:.2f}% of discharge")
