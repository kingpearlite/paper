"""Genset starts, run hours and loading in an exported dispatch (2026-09-18). Read-only.

Reads Time_Series_SC_*.xlsx of a Results folder: units running per hour = units at full load + units
in partial load (integral in a certified run). A start is any hourly increase in that count, carried
across year sheets. Prints, per scenario and overall, unit starts per year, unit-hours per year and
the mean loading of running units (genset energy / (unit-hours x unit size)).

    python _count_genset_starts.py <Results folder> [unit_kW=470]
"""
import glob
import os
import sys

import numpy as np
import openpyxl

folder = sys.argv[1]
unit_kw = float(sys.argv[2]) if len(sys.argv) > 2 else 470.0
files = sorted(glob.glob(os.path.join(folder, "Time_Series_SC_*.xlsx")),
               key=lambda p: int(p.rsplit("_", 1)[1].split(".")[0]))
tot_starts = tot_uh = tot_e = 0.0
n_years = 0
for f in files:
    wb = openpyxl.load_workbook(f, read_only=True)
    starts = uh = energy = 0.0
    prev = 0.0
    for name in wb.sheetnames:
        rows = [r for r in wb[name].iter_rows(min_row=5, values_only=True) if r[0] is not None]
        a = np.array([[float(x or 0) for x in r[1:8]] for r in rows])
        gen, n_part, n_full = a[:, 2], a[:, 4], a[:, 5]
        on = np.round(n_part + n_full)
        starts += np.clip(np.diff(np.concatenate([[prev], on])), 0, None).sum()
        prev = on[-1]
        uh += on.sum()
        energy += gen.sum()
    years = len(wb.sheetnames)
    wb.close()
    n_years = years
    print(f"{os.path.basename(f)}: {starts / years:6.0f} unit starts/yr, {uh / years:6.0f} unit-hours/yr, "
          f"mean loading {energy / uh / unit_kw if uh else 0:.0%}")
    tot_starts += starts
    tot_uh += uh
    tot_e += energy
k = len(files) * n_years
print(f"ALL {len(files)} scenario(s): {tot_starts / k:.0f} starts/yr, {tot_uh / k:.0f} unit-h/yr, "
      f"{tot_uh / tot_starts if tot_starts else 0:.1f} h per start, mean loading {tot_e / tot_uh / unit_kw:.0%}")
