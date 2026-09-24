"""Audit Eq. 5 on a certified run: residual, and how often the genset runs above net load
(so charges the battery or feeds curtailment) despite E_GEN <= E_dem."""
import sys, glob, os
import numpy as np
import openpyxl

folder = sys.argv[1]
files = sorted(glob.glob(os.path.join(folder, "Time_Series_SC_*.xlsx")),
               key=lambda p: int(p.rsplit("_", 1)[1].split(".")[0]))
tot = {}
print(f"{'sc':>3} {'dem GWh':>8} {'GEN':>7} {'RES':>7} {'ch':>6} {'curt':>6} {'LL':>5} {'maxres':>8} "
      f"{'hrs_gen>net':>11} {'E_gen>net MWh':>13} {'%ch_from_gen':>12} {'hrs_curt&gen':>12} {'hrs_gen=dem&res>0':>17}")
for f in files:
    sc = int(f.rsplit("_", 1)[1].split(".")[0])
    wb = openpyxl.load_workbook(f, read_only=True)
    acc = dict(dem=0, gen=0, res=0, ch=0, dis=0, curt=0, ll=0, h_ex=0, e_ex=0, ch_gen=0, h_cg=0, h_gd=0)
    maxres = 0.0
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(min_row=5, values_only=True))
        a = np.array([[float(x or 0) for x in r[1:12]] for r in rows if r[0] is not None])
        dem, res, gen, _, _, _, dis, ch, ll, curt, _ = a.T
        resid = res + gen + dis - ch - curt + ll - dem
        maxres = max(maxres, np.abs(resid).max())
        net = np.maximum(0, dem - res)
        ex = np.maximum(0, gen - net - dis)  # genset output beyond what load net of PV and discharge needs
        acc["dem"] += dem.sum(); acc["gen"] += gen.sum(); acc["res"] += res.sum()
        acc["ch"] += ch.sum(); acc["dis"] += dis.sum(); acc["curt"] += curt.sum(); acc["ll"] += ll.sum()
        m = ex > 1e-3
        acc["h_ex"] += m.sum(); acc["e_ex"] += ex.sum()
        acc["ch_gen"] += np.minimum(ch, ex).sum()
        acc["h_cg"] += ((curt > 1e-3) & (gen > 1e-3)).sum()
        acc["h_gd"] += ((gen >= dem - 1e-3) & (res > 1e-3) & (gen > 1e-3)).sum()
    wb.close()
    for k, v in acc.items():
        tot[k] = tot.get(k, 0) + v
    pct = 100 * acc["ch_gen"] / acc["ch"] if acc["ch"] else 0
    print(f"{sc:>3} {acc['dem']/1e6:8.2f} {acc['gen']/1e6:7.2f} {acc['res']/1e6:7.2f} {acc['ch']/1e6:6.2f} "
          f"{acc['curt']/1e6:6.2f} {acc['ll']:5.0f} {maxres:8.1e} {acc['h_ex']:11d} {acc['e_ex']/1e3:13.1f} "
          f"{pct:12.2f} {acc['h_cg']:12d} {acc['h_gd']:17d}")
pct = 100 * tot["ch_gen"] / tot["ch"]
print(f"ALL dem {tot['dem']/1e6:.2f} GWh, GEN {tot['gen']/1e6:.2f}, genset above net load {tot['e_ex']/1e3:.1f} MWh "
      f"({100*tot['e_ex']/tot['gen']:.3f}% of GEN) in {tot['h_ex']} of {len(files)*20*8760} hours; "
      f"{pct:.2f}% of battery charge; curtailment {tot['curt']/1e6:.3f} GWh, {tot['h_cg']} hours of curtailment with a genset on")
